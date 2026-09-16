from flask import (
    Flask,
    render_template,
    send_from_directory,
    request,
    redirect,
    url_for,
    session,
    abort,
    flash,
    Response,
    send_file,
)

import csv
import io
import os
import sqlite3
import sys
import time
import unicodedata
import math
from collections import Counter
from datetime import datetime, timedelta
from functools import wraps

from werkzeug.security import generate_password_hash, check_password_hash


BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from isg_database import (
    audit_kaydi_ekle as veritabanina_audit_kaydi_ekle,
    connect_database,
    prepare_database,
)
from isg_security import (
    generate_csrf_token,
    get_required_secret,
    login_attempt_key,
    validate_password,
    verify_csrf_token,
)


app = Flask(__name__)

app.secret_key = get_required_secret("ISG_SECRET_KEY")

app.secret_key = get_required_secret("ISG_SECRET_KEY")

app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=os.environ.get("ISG_COOKIE_SECURE", "0") == "1",
    PERMANENT_SESSION_LIFETIME=timedelta(hours=8),
)

CONTENT_SECURITY_POLICY = "; ".join(
    (
        "default-src 'self'",
        "base-uri 'self'",
        "object-src 'none'",
        "frame-ancestors 'none'",
        "form-action 'self'",
        "img-src 'self' data:",
        "font-src 'self'",
        "connect-src 'self'",
        "style-src 'self' 'unsafe-inline'",
        "script-src 'self' 'unsafe-inline'",
    )
)


# --------------------------------------------------
# SABİTLER
# --------------------------------------------------

MODULLER = [
    ("dashboard", "Dashboard"),
    ("ihlaller", "İhlaller"),
    ("kameralar", "Kameralar"),
    ("kullanicilar", "Kullanıcılar"),
    ("raporlar", "Raporlar"),
    ("ayarlar", "Sistem Ayarları"),
]

IZINLI_ROLLER = [
    "Admin",
    "İSG Uzmanı",
    "İSG Uzman Yardımcısı",
    "Görüntüleyici",
]

IZINLI_IHLAL_DURUMLARI = [
    "Yeni",
    "İnceleniyor",
    "Çözüldü",
]

ADMIN_ROLU = "Admin"

LOGIN_WINDOW_SECONDS = 15 * 60
LOGIN_BLOCK_SECONDS = 15 * 60
LOGIN_MAX_FAILURES = 5

KULLANICI_ADI_MIN_UZUNLUK = 2
KULLANICI_ADI_MAX_UZUNLUK = 64
AD_SOYAD_MIN_UZUNLUK = 2
AD_SOYAD_MAX_UZUNLUK = 120
PAROLA_MAX_UZUNLUK = 128
ACIKLAMA_MAX_UZUNLUK = 3000
KAMERA_ADI_MIN_UZUNLUK = 2
KAMERA_ADI_MAX_UZUNLUK = 100
IHLAL_TURU_MAX_UZUNLUK = 200
GENEL_LOGIN_HATASI = (
    "Kullanıcı adı veya parola hatalı ya da hesap kullanılamıyor."
)


def kontrol_karakteri_iceriyor(deger):
    return any(unicodedata.category(karakter).startswith("C") for karakter in deger)


def metin_gecerli(deger, minimum, maksimum, bos_olabilir=False):
    if bos_olabilir and deger == "":
        return True

    return (
        minimum <= len(deger) <= maksimum
        and not kontrol_karakteri_iceriyor(deger)
    )


def kamera_formunu_dogrula(form):
    kamera_adi = form.get("kamera_adi", "").strip()
    kamera_index_degeri = form.get("kamera_index", "").strip()
    yasak_bolge_yuzdesi_degeri = form.get(
        "yasak_bolge_yuzdesi", ""
    ).strip()

    if not metin_gecerli(
        kamera_adi,
        KAMERA_ADI_MIN_UZUNLUK,
        KAMERA_ADI_MAX_UZUNLUK,
    ):
        return None, "Kamera adı geçersiz."

    try:
        kamera_index = int(kamera_index_degeri)
    except ValueError:
        return None, "Kamera index değeri geçersiz."

    if kamera_index < 0:
        return None, "Kamera index değeri negatif olamaz."

    try:
        yasak_bolge_yuzdesi = float(yasak_bolge_yuzdesi_degeri)
    except ValueError:
        return None, "Yasak bölge yüzdesi geçersiz."

    if (
        not math.isfinite(yasak_bolge_yuzdesi)
        or yasak_bolge_yuzdesi < 0
        or yasak_bolge_yuzdesi > 100
    ):
        return None, "Yasak bölge yüzdesi 0 ile 100 arasında olmalıdır."

    return {
        "kamera_adi": kamera_adi,
        "kamera_index": kamera_index,
        "yasak_bolge_orani": yasak_bolge_yuzdesi / 100,
    }, None


def rapor_filtrelerini_dogrula(args, izinli_ihlal_turleri):
    baslangic = args.get("baslangic", "").strip()
    bitis = args.get("bitis", "").strip()
    durum = args.get("durum", "").strip()
    ihlal_turu = args.get("ihlal_turu", "").strip()

    try:
        baslangic_tarihi = (
            datetime.strptime(baslangic, "%Y-%m-%d").date()
            if baslangic
            else None
        )
        bitis_tarihi = (
            datetime.strptime(bitis, "%Y-%m-%d").date()
            if bitis
            else None
        )
    except ValueError:
        return None, "Tarih formatı YYYY-MM-DD olmalıdır."

    if baslangic_tarihi and bitis_tarihi and baslangic_tarihi > bitis_tarihi:
        return None, "Başlangıç tarihi bitiş tarihinden büyük olamaz."

    if durum and durum not in IZINLI_IHLAL_DURUMLARI:
        return None, "Geçersiz ihlal durumu."

    if ihlal_turu:
        if not metin_gecerli(ihlal_turu, 1, IHLAL_TURU_MAX_UZUNLUK):
            return None, "Geçersiz ihlal türü."

        if ihlal_turu not in izinli_ihlal_turleri:
            return None, "Geçersiz ihlal türü."

    kamera_id = args.get("kamera_id", "").strip()
    if kamera_id and (not kamera_id.isascii() or not kamera_id.isdigit() or not 0 < int(kamera_id) <= 2147483647):
        return None, "Geçersiz kamera."
    q = args.get("q", "").strip()
    if len(q) > 120 or kontrol_karakteri_iceriyor(q):
        return None, "Arama en fazla 120 karakter olmalıdır."

    return {
        "kamera_id": kamera_id,
        "q": q,
        "baslangic": baslangic,
        "bitis": bitis,
        "durum": durum,
        "ihlal_turu": ihlal_turu,
    }, None


def rapor_where_clause(filtreler):
    kosullar = []
    parametreler = []

    if filtreler["baslangic"]:
        kosullar.append("tarih >= ?")
        parametreler.append(f"{filtreler['baslangic']} 00:00:00")

    if filtreler["bitis"]:
        bitis_siniri = datetime.strptime(
            filtreler["bitis"], "%Y-%m-%d"
        ) + timedelta(days=1)
        kosullar.append("tarih < ?")
        parametreler.append(bitis_siniri.strftime("%Y-%m-%d 00:00:00"))

    if filtreler["durum"]:
        kosullar.append("durum = ?")
        parametreler.append(filtreler["durum"])

    if filtreler["ihlal_turu"]:
        kosullar.append("ihlal_turu = ?")
        parametreler.append(filtreler["ihlal_turu"])

    if filtreler.get("kamera_id"):
        kosullar.append("kamera_id = ?")
        parametreler.append(int(filtreler["kamera_id"]))
    if filtreler.get("q"):
        kosullar.append("(CAST(id AS TEXT) = ? OR instr(COALESCE(person_id, ''), ?) > 0 OR instr(COALESCE(aciklama, ''), ?) > 0)")
        parametreler.extend([filtreler["q"]] * 3)

    where_sql = " WHERE " + " AND ".join(kosullar) if kosullar else ""
    return where_sql, parametreler


def tablo_hucresini_guvenli_yap(deger):
    metin = "" if deger is None else str(deger)
    kontrol_metni = metin.lstrip()

    if kontrol_metni and kontrol_metni[0] in "=+-@":
        return "'" + metin

    return metin


def csv_hucresini_guvenli_yap(deger):
    return tablo_hucresini_guvenli_yap(deger)


# --------------------------------------------------
# VERİTABANI HAZIRLAMA
# --------------------------------------------------

def kullanici_tablosunu_hazirla():
    baglanti = connect_database()
    cursor = baglanti.cursor()

    cursor.execute(
        """
        SELECT id, sifre
        FROM kullanicilar
        WHERE kullanici_adi = ?
        """,
        ("admin",),
    )

    admin = cursor.fetchone()

    bootstrap_password = os.environ.pop(
        "ISG_BOOTSTRAP_ADMIN_PASSWORD",
        "",
    )

    if admin is None:
        is_valid, error_message = validate_password(bootstrap_password)

        if not is_valid:
            baglanti.close()
            raise RuntimeError(
                "İlk admin hesabını oluşturmak için güçlü bir "
                "ISG_BOOTSTRAP_ADMIN_PASSWORD tanımlanmalıdır. "
                + error_message
            )

        sifre_hash = generate_password_hash(bootstrap_password)

        cursor.execute(
            """
            INSERT INTO kullanicilar
                (ad_soyad, kullanici_adi, sifre, rol, aktif)
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                "Sistem Yöneticisi",
                "admin",
                sifre_hash,
                "Admin",
                1,
            ),
        )

    elif check_password_hash(admin[1], "admin123"):
        is_valid, error_message = validate_password(bootstrap_password)

        if not is_valid:
            baglanti.close()
            raise RuntimeError(
                "Mevcut admin123 parolası devre dışı bırakılmalıdır. "
                "Güçlü bir ISG_BOOTSTRAP_ADMIN_PASSWORD tanımlayıp "
                "uygulamayı yeniden başlatın. "
                + error_message
            )

        cursor.execute(
            """
            UPDATE kullanicilar
            SET sifre = ?
            WHERE id = ?
            """,
            (generate_password_hash(bootstrap_password), admin[0]),
        )

    baglanti.commit()
    baglanti.close()


def yetki_tablosunu_hazirla():
    baglanti = connect_database()
    cursor = baglanti.cursor()

    # İlk admin hesabı sistemden kilitlenmesin diye tüm modülleri veriyoruz.
    cursor.execute(
        """
        SELECT id
        FROM kullanicilar
        WHERE kullanici_adi = ?
        """,
        ("admin",),
    )

    admin = cursor.fetchone()

    if admin:
        admin_id = admin[0]

        for modul_kodu, _ in MODULLER:
            cursor.execute(
                """
                INSERT OR IGNORE INTO kullanici_yetkileri
                    (kullanici_id, modul)
                VALUES (?, ?)
                """,
                (admin_id, modul_kodu),
            )

    baglanti.commit()
    baglanti.close()


def veritabanini_hazirla():
    prepare_database()
    kullanici_tablosunu_hazirla()
    yetki_tablosunu_hazirla()


veritabanini_hazirla()


# --------------------------------------------------
# İSTEK / SESSION GÜVENLİĞİ
# --------------------------------------------------

@app.context_processor
def csrf_context():
    return {
        "csrf_token": lambda: generate_csrf_token(session),
        "admin_mi": lambda: session.get("rol") == ADMIN_ROLU,
        "input_limitleri": {
            "kullanici_adi_min": KULLANICI_ADI_MIN_UZUNLUK,
            "kullanici_adi_max": KULLANICI_ADI_MAX_UZUNLUK,
            "ad_soyad_min": AD_SOYAD_MIN_UZUNLUK,
            "ad_soyad_max": AD_SOYAD_MAX_UZUNLUK,
            "parola_max": PAROLA_MAX_UZUNLUK,
            "aciklama_max": ACIKLAMA_MAX_UZUNLUK,
            "kamera_adi_min": KAMERA_ADI_MIN_UZUNLUK,
            "kamera_adi_max": KAMERA_ADI_MAX_UZUNLUK,
        },
    }


@app.after_request
def guvenlik_basliklarini_ekle(response):
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "strict-origin-when-cross-origin"
    response.headers["Content-Security-Policy"] = CONTENT_SECURITY_POLICY

    if request.is_secure or app.config["SESSION_COOKIE_SECURE"]:
        response.headers["Strict-Transport-Security"] = (
            "max-age=31536000; includeSubDomains"
        )

    return response


@app.before_request
def guvenlik_kontrolleri():
    if request.method == "POST":
        csrf_candidate = (
            request.form.get("csrf_token", "")
            or request.headers.get("X-CSRF-Token", "")
        )

        if not verify_csrf_token(session, csrf_candidate):
            abort(400, description="Geçersiz veya eksik güvenlik doğrulaması.")

    kullanici_id = session.get("kullanici_id")

    if not kullanici_id:
        return None

    baglanti = connect_database()
    cursor = baglanti.cursor()
    cursor.execute(
        """
        SELECT ad_soyad, kullanici_adi, rol, aktif
        FROM kullanicilar
        WHERE id = ?
        """,
        (kullanici_id,),
    )
    kullanici = cursor.fetchone()

    if kullanici is None or kullanici[3] != 1:
        baglanti.close()
        session.clear()
        return redirect(url_for("login"))

    cursor.execute(
        """
        SELECT modul
        FROM kullanici_yetkileri
        WHERE kullanici_id = ?
        ORDER BY id
        """,
        (kullanici_id,),
    )
    yetkiler = [kayit[0] for kayit in cursor.fetchall()]
    baglanti.close()

    if not yetkiler:
        session.clear()
        return redirect(url_for("login"))

    session["ad_soyad"] = kullanici[0]
    session["kullanici_adi"] = kullanici[1]
    session["rol"] = kullanici[2]
    session["yetkiler"] = yetkiler

    return None


# --------------------------------------------------
# YETKİ YARDIMCILARI
# --------------------------------------------------

def kullanici_yetkilerini_getir(kullanici_id):
    baglanti = connect_database()
    cursor = baglanti.cursor()

    cursor.execute(
        """
        SELECT modul
        FROM kullanici_yetkileri
        WHERE kullanici_id = ?
        ORDER BY id
        """,
        (kullanici_id,),
    )

    yetkiler = [kayit[0] for kayit in cursor.fetchall()]

    baglanti.close()
    return yetkiler


def login_gerekli(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if "kullanici_id" not in session:
            return redirect(url_for("login"))

        return f(*args, **kwargs)

    return decorated_function


def yetki_gerekli(modul):
    def decorator(f):
        @wraps(f)
        def decorated_function(*args, **kwargs):
            if "kullanici_id" not in session:
                return redirect(url_for("login"))

            yetkiler = session.get("yetkiler", [])

            if modul not in yetkiler:
                return "Bu modüle erişim yetkiniz bulunmuyor.", 403

            return f(*args, **kwargs)

        return decorated_function

    return decorator


def admin_gerekli(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if "kullanici_id" not in session:
            return redirect(url_for("login"))

        if session.get("rol") != ADMIN_ROLU:
            abort(403, description="Bu işlem için yönetici yetkisi gereklidir.")

        return f(*args, **kwargs)

    return decorated_function


def aktif_admin_sayisi(connection):
    return connection.execute(
        """
        SELECT COUNT(*)
        FROM kullanicilar
        WHERE rol = ? AND aktif = 1
        """,
        (ADMIN_ROLU,),
    ).fetchone()[0]


def audit_kaydi_ekle(
    connection,
    islem_turu,
    hedef_turu,
    hedef_id,
    eski_deger=None,
    yeni_deger=None,
):
    veritabanina_audit_kaydi_ekle(
        connection=connection,
        kullanici_id=session.get("kullanici_id"),
        kullanici_adi=session.get("kullanici_adi", "bilinmiyor"),
        islem_turu=islem_turu,
        hedef_turu=hedef_turu,
        hedef_id=hedef_id,
        eski_deger=eski_deger,
        yeni_deger=yeni_deger,
        ip_adresi=request.remote_addr or "unknown",
    )


def ilk_erisilebilir_sayfa():
    yetkiler = session.get("yetkiler", [])

    if not yetkiler:
        return url_for("login")

    if "dashboard" in yetkiler:
        return url_for("ana_sayfa")

    if "ihlaller" in yetkiler:
        return url_for("ihlal_listesi")

    if "raporlar" in yetkiler:
        return url_for("raporlar")

    if "kameralar" in yetkiler:
        return url_for("kameralar")

    if "kullanicilar" in yetkiler:
        return url_for("kullanicilar")

    # Diğer modüllerin ayrı route'ları eklendikçe buraya bağlanabilir.
    return url_for("yetkisiz")


# --------------------------------------------------
# LOGIN / LOGOUT
# --------------------------------------------------

def login_engel_suresi(anahtar):
    simdi = int(time.time())
    baglanti = connect_database()
    cursor = baglanti.cursor()
    cursor.execute(
        """
        SELECT pencere_baslangici, engel_bitis
        FROM giris_denemeleri
        WHERE anahtar = ?
        """,
        (anahtar,),
    )
    kayit = cursor.fetchone()

    if kayit is None:
        baglanti.close()
        return 0

    if kayit[1] > simdi:
        baglanti.close()
        return kayit[1] - simdi

    if simdi - kayit[0] >= LOGIN_WINDOW_SECONDS:
        cursor.execute(
            "DELETE FROM giris_denemeleri WHERE anahtar = ?",
            (anahtar,),
        )
        baglanti.commit()

    baglanti.close()
    return 0


def basarisiz_girisi_kaydet(anahtar):
    simdi = int(time.time())
    baglanti = connect_database()
    cursor = baglanti.cursor()
    cursor.execute(
        """
        SELECT basarisiz_sayisi, pencere_baslangici
        FROM giris_denemeleri
        WHERE anahtar = ?
        """,
        (anahtar,),
    )
    kayit = cursor.fetchone()

    if kayit is None or simdi - kayit[1] >= LOGIN_WINDOW_SECONDS:
        basarisiz_sayisi = 1
        pencere_baslangici = simdi
    else:
        basarisiz_sayisi = kayit[0] + 1
        pencere_baslangici = kayit[1]

    engel_bitis = 0

    if basarisiz_sayisi >= LOGIN_MAX_FAILURES:
        engel_bitis = simdi + LOGIN_BLOCK_SECONDS

    cursor.execute(
        """
        INSERT INTO giris_denemeleri
            (anahtar, basarisiz_sayisi, pencere_baslangici, engel_bitis)
        VALUES (?, ?, ?, ?)
        ON CONFLICT(anahtar) DO UPDATE SET
            basarisiz_sayisi = excluded.basarisiz_sayisi,
            pencere_baslangici = excluded.pencere_baslangici,
            engel_bitis = excluded.engel_bitis
        """,
        (
            anahtar,
            basarisiz_sayisi,
            pencere_baslangici,
            engel_bitis,
        ),
    )
    baglanti.commit()
    baglanti.close()


def giris_denemelerini_temizle(anahtar):
    baglanti = connect_database()
    baglanti.execute(
        "DELETE FROM giris_denemeleri WHERE anahtar = ?",
        (anahtar,),
    )
    baglanti.commit()
    baglanti.close()

@app.route("/login", methods=["GET", "POST"])
def login():
    hata = None

    if "kullanici_id" in session:
        return redirect(ilk_erisilebilir_sayfa())

    if request.method == "POST":
        kullanici_adi = request.form.get("kullanici_adi", "").strip()
        sifre = request.form.get("sifre", "")
        deneme_anahtari = login_attempt_key(
            request.remote_addr or "unknown",
            kullanici_adi,
        )
        kalan_sure = login_engel_suresi(deneme_anahtari)

        if kalan_sure:
            kalan_dakika = max(1, (kalan_sure + 59) // 60)
            hata = (
                "Çok fazla başarısız giriş denemesi yapıldı. "
                f"Yaklaşık {kalan_dakika} dakika sonra tekrar deneyin."
            )
            return render_template("login.html", hata=hata), 429

        login_girdisi_gecerli = (
            metin_gecerli(
                kullanici_adi,
                KULLANICI_ADI_MIN_UZUNLUK,
                KULLANICI_ADI_MAX_UZUNLUK,
            )
            and 0 < len(sifre) <= PAROLA_MAX_UZUNLUK
        )

        kullanici = None

        if login_girdisi_gecerli:
            baglanti = connect_database()
            kullanici = baglanti.execute(
                """
                SELECT
                    id,
                    ad_soyad,
                    kullanici_adi,
                    sifre,
                    rol,
                    aktif
                FROM kullanicilar
                WHERE kullanici_adi = ?
                """,
                (kullanici_adi,),
            ).fetchone()
            baglanti.close()

        if kullanici is None:
            hata = GENEL_LOGIN_HATASI

        elif kullanici[5] != 1:
            hata = GENEL_LOGIN_HATASI

        elif not check_password_hash(kullanici[3], sifre):
            hata = GENEL_LOGIN_HATASI

        else:
            giris_denemelerini_temizle(deneme_anahtari)
            session.clear()
            session.permanent = True

            session["kullanici_id"] = kullanici[0]
            session["ad_soyad"] = kullanici[1]
            session["kullanici_adi"] = kullanici[2]
            session["rol"] = kullanici[4]
            session["yetkiler"] = kullanici_yetkilerini_getir(kullanici[0])

            if not session["yetkiler"]:
                session.clear()
                hata = GENEL_LOGIN_HATASI
            else:
                return redirect(ilk_erisilebilir_sayfa())

        if hata:
            basarisiz_girisi_kaydet(deneme_anahtari)

    return render_template("login.html", hata=hata)


@app.route("/logout", methods=["POST"])
@login_gerekli
def logout():
    session.clear()
    return redirect(url_for("login"))


@app.route("/yetkisiz")
@login_gerekli
def yetkisiz():
    return (
        "Bu kullanıcıya henüz erişebileceği bir modül atanmadı. "
        "Lütfen sistem yöneticisine başvurun.",
        403,
    )


# --------------------------------------------------
# DASHBOARD
# --------------------------------------------------

@app.route("/")
def kok_sayfa():
    if "kullanici_id" not in session:
        return redirect(url_for("login"))

    yetkiler = session.get("yetkiler", [])

    if not yetkiler:
        session.clear()
        return redirect(url_for("login"))

    return redirect(ilk_erisilebilir_sayfa())


@app.route("/dashboard")
@yetki_gerekli("dashboard")
def ana_sayfa():
    baglanti = connect_database()
    cursor = baglanti.cursor()

    cursor.execute(
        """
        SELECT
            id,
            tarih,
            ihlal_turu,
            guven_skoru,
            fotograf_yolu,
            durum
        FROM ihlaller
        ORDER BY id DESC
        """
    )

    ihlaller = cursor.fetchall()
    baglanti.close()

    return render_template("index.html", ihlaller=ihlaller)


# --------------------------------------------------
# İHLAL İŞLEMLERİ
# --------------------------------------------------

@app.route("/durum-guncelle/<int:ihlal_id>", methods=["POST"])
@yetki_gerekli("ihlaller")
def durum_guncelle(ihlal_id):
    if session.get("rol") == "Görüntüleyici":
        abort(403)
    yeni_durum = request.form.get("durum")
    cozum_notu = request.form.get("cozum_notu", "").strip()
    if len(cozum_notu) > ACIKLAMA_MAX_UZUNLUK:
        abort(400, description="Çözüm notu çok uzun.")

    if yeni_durum not in IZINLI_IHLAL_DURUMLARI:
        return "Geçersiz durum", 400

    baglanti = connect_database()

    try:
        baglanti.execute("BEGIN IMMEDIATE")
        ihlal = baglanti.execute(
            "SELECT durum, cozum_notu, cozen_kullanici_id, cozum_tarihi FROM ihlaller WHERE id = ?",
            (ihlal_id,),
        ).fetchone()

        if ihlal is None:
            baglanti.rollback()
            return "İhlal bulunamadı", 404

        # Repeated submissions do not overwrite the original resolver/time.
        if ihlal[0] == yeni_durum:
            baglanti.rollback()
            return redirect(url_for("ihlal_detay", ihlal_id=ihlal_id))
        resolved = yeni_durum == "Çözüldü"
        cozum_tarihi = datetime.now().isoformat(timespec="seconds") if resolved else None
        baglanti.execute(
            "UPDATE ihlaller SET durum=?, cozum_notu=?, cozen_kullanici_id=?, cozen_kullanici_adi=?, cozum_tarihi=? WHERE id=?",
            (yeni_durum, cozum_notu if resolved else None,
             session["kullanici_id"] if resolved else None,
             session.get("ad_soyad") if resolved else None, cozum_tarihi, ihlal_id),
        )
        audit_kaydi_ekle(
            baglanti,
            "ihlal_durumu_degistirildi",
            "ihlal",
            ihlal_id,
            eski_deger={"durum": ihlal[0], "cozum_notu": ihlal[1], "cozen_kullanici_id": ihlal[2], "cozum_tarihi": ihlal[3]},
            yeni_deger={"durum": yeni_durum, "cozum_notu": cozum_notu if resolved else None, "cozum_tarihi": cozum_tarihi},
        )
        baglanti.commit()
    except Exception:
        baglanti.rollback()
        raise
    finally:
        baglanti.close()

    return redirect(url_for("ihlal_detay", ihlal_id=ihlal_id))


@app.route("/ihlal/<int:ihlal_id>")
@yetki_gerekli("ihlaller")
def ihlal_detay(ihlal_id):
    baglanti = connect_database()
    cursor = baglanti.cursor()

    cursor.execute(
        """
        SELECT
            id,
            tarih,
            ihlal_turu,
            guven_skoru,
            fotograf_yolu,
            durum,
            aciklama
        FROM ihlaller
        WHERE id = ?
        """,
        (ihlal_id,),
    )

    ihlal = cursor.fetchone()
    baglanti.close()

    if ihlal is None:
        return "İhlal bulunamadı", 404

    return render_template("detay.html", ihlal=ihlal)


@app.route("/ihlal/<int:ihlal_id>/aciklama", methods=["POST"])
@yetki_gerekli("ihlaller")
def aciklama_kaydet(ihlal_id):
    if session.get("rol") == "Görüntüleyici":
        abort(403)
    aciklama = request.form.get("aciklama", "").strip()

    if len(aciklama) > ACIKLAMA_MAX_UZUNLUK:
        return "Açıklama izin verilen uzunluğu aşıyor.", 400

    baglanti = connect_database()

    try:
        baglanti.execute("BEGIN IMMEDIATE")
        ihlal = baglanti.execute(
            "SELECT aciklama FROM ihlaller WHERE id = ?",
            (ihlal_id,),
        ).fetchone()

        if ihlal is None:
            baglanti.rollback()
            return "İhlal bulunamadı", 404

        baglanti.execute(
            "UPDATE ihlaller SET aciklama = ? WHERE id = ?",
            (aciklama, ihlal_id),
        )
        audit_kaydi_ekle(
            baglanti,
            "ihlal_aciklamasi_degistirildi",
            "ihlal",
            ihlal_id,
            eski_deger={"aciklama": ihlal[0]},
            yeni_deger={"aciklama": aciklama},
        )
        baglanti.commit()
    except Exception:
        baglanti.rollback()
        raise
    finally:
        baglanti.close()

    return redirect(url_for("ihlal_detay", ihlal_id=ihlal_id))


@app.route("/ihlaller/<path:dosya_adi>")
@yetki_gerekli("ihlaller")
def ihlal_gorseli(dosya_adi):
    proje_klasoru = os.path.dirname(app.root_path)
    klasor = os.path.join(proje_klasoru, "ihlaller")

    return send_from_directory(klasor, dosya_adi)


# --------------------------------------------------
# RAPORLAR
# --------------------------------------------------

def ihlal_turlerini_getir(baglanti):
    return [
        row[0]
        for row in baglanti.execute(
            """
            SELECT DISTINCT ihlal_turu
            FROM ihlaller
            WHERE ihlal_turu IS NOT NULL AND ihlal_turu != ''
            ORDER BY ihlal_turu
            """
        ).fetchall()
    ]


def rapor_verilerini_getir(baglanti, args, aciklama_dahil=False):
    ihlal_turleri = ihlal_turlerini_getir(baglanti)
    filtreler, hata = rapor_filtrelerini_dogrula(args, ihlal_turleri)

    if hata:
        return ihlal_turleri, None, None, hata

    where_sql, parametreler = rapor_where_clause(filtreler)
    kolonlar = "id, tarih, ihlal_turu, guven_skoru, durum"

    if aciklama_dahil:
        kolonlar += ", aciklama"

    ihlaller = baglanti.execute(
        f"SELECT {kolonlar} FROM ihlaller"
        + where_sql
        + " ORDER BY tarih DESC, id DESC",
        parametreler,
    ).fetchall()

    return ihlal_turleri, filtreler, ihlaller, None


class XlsxBagimliligiHatasi(RuntimeError):
    pass


class PdfBagimliligiHatasi(RuntimeError):
    pass


class PdfFontHatasi(RuntimeError):
    pass


def xlsx_workbook_olustur(ihlaller, filtreler):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError as exc:
        raise XlsxBagimliligiHatasi(
            "XLSX dışa aktarma bağımlılığı kullanılamıyor."
        ) from exc

    workbook = Workbook()
    rapor_sayfasi = workbook.active
    rapor_sayfasi.title = "İhlal Raporu"
    ozet_sayfasi = workbook.create_sheet("Özet")

    baslik_dolgusu = PatternFill("solid", fgColor="1F4E78")
    alt_baslik_dolgusu = PatternFill("solid", fgColor="D9EAF7")
    beyaz_kalin_yazi = Font(color="FFFFFF", bold=True)
    kalin_yazi = Font(bold=True)
    ince_kenarlik = Border(
        bottom=Side(style="thin", color="B8C6D1")
    )

    rapor_sayfasi.merge_cells("A1:F1")
    rapor_sayfasi["A1"] = "İSG İhlal Raporu"
    rapor_sayfasi["A1"].fill = baslik_dolgusu
    rapor_sayfasi["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    rapor_sayfasi["A1"].alignment = Alignment(horizontal="center")

    tarih_araligi = "Tümü"
    if filtreler["baslangic"] or filtreler["bitis"]:
        tarih_araligi = (
            f"{filtreler['baslangic'] or 'Başlangıç yok'} - "
            f"{filtreler['bitis'] or 'Bitiş yok'}"
        )

    rapor_bilgileri = [
        ("Rapor oluşturma tarihi", datetime.now()),
        ("Tarih aralığı", tarih_araligi),
        ("Durum filtresi", filtreler["durum"] or "Tümü"),
        ("İhlal türü filtresi", filtreler["ihlal_turu"] or "Tümü"),
        ("Toplam kayıt", len(ihlaller)),
    ]

    for satir, (etiket, deger) in enumerate(rapor_bilgileri, start=3):
        rapor_sayfasi.cell(satir, 1, etiket).font = kalin_yazi
        rapor_sayfasi.cell(satir, 2, deger)

    rapor_sayfasi["B3"].number_format = "yyyy-mm-dd hh:mm:ss"
    tablo_baslik_satiri = 9
    kolon_basliklari = [
        "ID",
        "Tarih",
        "İhlal Türü",
        "Güven Skoru (%)",
        "Durum",
        "Açıklama",
    ]

    for kolon, baslik in enumerate(kolon_basliklari, start=1):
        hucre = rapor_sayfasi.cell(tablo_baslik_satiri, kolon, baslik)
        hucre.fill = baslik_dolgusu
        hucre.font = beyaz_kalin_yazi
        hucre.alignment = Alignment(horizontal="center")

    for satir, ihlal in enumerate(ihlaller, start=tablo_baslik_satiri + 1):
        tarih_degeri = ihlal[1]
        try:
            tarih_degeri = datetime.fromisoformat(str(tarih_degeri))
        except (TypeError, ValueError):
            tarih_degeri = tablo_hucresini_guvenli_yap(tarih_degeri)

        guven_skoru = None
        if ihlal[3] is not None:
            guven_skoru = float(ihlal[3]) / 100

        degerler = [
            ihlal[0],
            tarih_degeri,
            tablo_hucresini_guvenli_yap(ihlal[2]),
            guven_skoru,
            tablo_hucresini_guvenli_yap(ihlal[4]),
            tablo_hucresini_guvenli_yap(ihlal[5]),
        ]

        for kolon, deger in enumerate(degerler, start=1):
            hucre = rapor_sayfasi.cell(satir, kolon, deger)
            hucre.border = ince_kenarlik

        rapor_sayfasi.cell(satir, 2).number_format = "yyyy-mm-dd hh:mm:ss"
        rapor_sayfasi.cell(satir, 4).number_format = "0.00%"
        rapor_sayfasi.cell(satir, 6).alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

    son_satir = max(tablo_baslik_satiri, tablo_baslik_satiri + len(ihlaller))
    rapor_sayfasi.freeze_panes = "A10"
    rapor_sayfasi.auto_filter.ref = f"A{tablo_baslik_satiri}:F{son_satir}"
    rapor_sayfasi.sheet_view.showGridLines = False
    kolon_genislikleri = {
        "A": 10,
        "B": 21,
        "C": 30,
        "D": 19,
        "E": 18,
        "F": 55,
    }
    for kolon, genislik in kolon_genislikleri.items():
        rapor_sayfasi.column_dimensions[kolon].width = genislik

    durum_sayilari = Counter(row[4] for row in ihlaller)
    gunluk_sayilar = Counter(str(row[1])[:10] for row in ihlaller)

    ozet_sayfasi.merge_cells("A1:B1")
    ozet_sayfasi["A1"] = "Rapor Özeti"
    ozet_sayfasi["A1"].fill = baslik_dolgusu
    ozet_sayfasi["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    ozet_sayfasi["A1"].alignment = Alignment(horizontal="center")

    ozet_degerleri = [
        ("Toplam ihlal", len(ihlaller)),
        ("Yeni", durum_sayilari.get("Yeni", 0)),
        ("İnceleniyor", durum_sayilari.get("İnceleniyor", 0)),
        ("Çözüldü", durum_sayilari.get("Çözüldü", 0)),
    ]
    for satir, (etiket, deger) in enumerate(ozet_degerleri, start=3):
        ozet_sayfasi.cell(satir, 1, etiket).font = kalin_yazi
        ozet_sayfasi.cell(satir, 2, deger)

    dagilim_baslik_satiri = 9
    for kolon, baslik in enumerate(("Tarih", "İhlal Sayısı"), start=1):
        hucre = ozet_sayfasi.cell(dagilim_baslik_satiri, kolon, baslik)
        hucre.fill = alt_baslik_dolgusu
        hucre.font = kalin_yazi

    for satir, (tarih, adet) in enumerate(
        sorted(gunluk_sayilar.items()),
        start=dagilim_baslik_satiri + 1,
    ):
        ozet_sayfasi.cell(satir, 1, tarih)
        ozet_sayfasi.cell(satir, 2, adet)

    ozet_sayfasi.freeze_panes = "A10"
    ozet_sayfasi.column_dimensions["A"].width = 24
    ozet_sayfasi.column_dimensions["B"].width = 16
    ozet_sayfasi.sheet_view.showGridLines = False

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def pdf_fontlarini_kaydet(pdfmetrics, TTFont):
    font_secenekleri = [
        (
            "C:/Windows/Fonts/arial.ttf",
            "C:/Windows/Fonts/arialbd.ttf",
        ),
        (
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        ),
        (
            "/System/Library/Fonts/Supplemental/Arial.ttf",
            "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
        ),
    ]

    for normal_font, kalin_font in font_secenekleri:
        if os.path.isfile(normal_font) and os.path.isfile(kalin_font):
            if "IsgUnicode" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("IsgUnicode", normal_font))
            if "IsgUnicode-Bold" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("IsgUnicode-Bold", kalin_font))
            return "IsgUnicode", "IsgUnicode-Bold"

    raise PdfFontHatasi(
        "Türkçe karakterleri destekleyen sistem fontu bulunamadı."
    )


def pdf_raporu_olustur(ihlaller, filtreler):
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER, TA_LEFT
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import (
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )
        from xml.sax.saxutils import escape
    except ImportError as exc:
        raise PdfBagimliligiHatasi(
            "PDF dışa aktarma bağımlılığı kullanılamıyor."
        ) from exc

    normal_font, kalin_font = pdf_fontlarini_kaydet(pdfmetrics, TTFont)
    output = io.BytesIO()
    sayfa_boyutu = landscape(A4)
    belge = SimpleDocTemplate(
        output,
        pagesize=sayfa_boyutu,
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=12 * mm,
        bottomMargin=13 * mm,
        title="İSG İhlal Raporu",
        author="İSG AI Sistemi",
    )

    stiller = getSampleStyleSheet()
    baslik_stili = ParagraphStyle(
        "IsgBaslik",
        parent=stiller["Title"],
        fontName=kalin_font,
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#163A5F"),
        alignment=TA_CENTER,
        spaceAfter=6 * mm,
    )
    bolum_stili = ParagraphStyle(
        "IsgBolum",
        parent=stiller["Heading2"],
        fontName=kalin_font,
        fontSize=11,
        leading=14,
        textColor=colors.HexColor("#163A5F"),
        spaceBefore=3 * mm,
        spaceAfter=2 * mm,
    )
    govde_stili = ParagraphStyle(
        "IsgGovde",
        parent=stiller["BodyText"],
        fontName=normal_font,
        fontSize=7.5,
        leading=10,
        alignment=TA_LEFT,
        wordWrap="LTR",
    )
    kalin_govde_stili = ParagraphStyle(
        "IsgKalinGovde",
        parent=govde_stili,
        fontName=kalin_font,
    )
    tablo_baslik_stili = ParagraphStyle(
        "IsgTabloBaslik",
        parent=govde_stili,
        fontName=kalin_font,
        textColor=colors.white,
        alignment=TA_CENTER,
    )

    def paragraf(deger, stil=govde_stili):
        metin = "" if deger is None else str(deger)
        return Paragraph(escape(metin).replace("\n", "<br/>"), stil)

    ana_mavi = colors.HexColor("#1F4E78")
    acik_mavi = colors.HexColor("#D9EAF7")
    cizgi_rengi = colors.HexColor("#B8C6D1")
    hikaye = [Paragraph("İSG İhlal Raporu", baslik_stili)]

    rapor_bilgileri = [
        ("Rapor oluşturma tarihi", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("Başlangıç tarihi", filtreler["baslangic"] or "Tümü"),
        ("Bitiş tarihi", filtreler["bitis"] or "Tümü"),
        ("Durum filtresi", filtreler["durum"] or "Tümü"),
        ("İhlal türü filtresi", filtreler["ihlal_turu"] or "Tümü"),
        ("Toplam kayıt", len(ihlaller)),
    ]
    bilgi_verileri = []
    for indeks in range(0, len(rapor_bilgileri), 2):
        satir = []
        for etiket, deger in rapor_bilgileri[indeks:indeks + 2]:
            satir.extend(
                [
                    paragraf(etiket, kalin_govde_stili),
                    paragraf(deger),
                ]
            )
        bilgi_verileri.append(satir)

    bilgi_tablosu = Table(
        bilgi_verileri,
        colWidths=[42 * mm, 52 * mm, 42 * mm, 80 * mm],
        hAlign="LEFT",
    )
    bilgi_tablosu.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), acik_mavi),
                ("BACKGROUND", (2, 0), (2, -1), acik_mavi),
                ("BOX", (0, 0), (-1, -1), 0.5, cizgi_rengi),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, cizgi_rengi),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    hikaye.extend([bilgi_tablosu, Spacer(1, 4 * mm)])

    durum_sayilari = Counter(row[4] for row in ihlaller)
    gunluk_sayilar = Counter(str(row[1])[:10] for row in ihlaller)
    ozet_verileri = [
        [
            paragraf("Toplam", tablo_baslik_stili),
            paragraf("Yeni", tablo_baslik_stili),
            paragraf("İnceleniyor", tablo_baslik_stili),
            paragraf("Çözüldü", tablo_baslik_stili),
        ],
        [
            len(ihlaller),
            durum_sayilari.get("Yeni", 0),
            durum_sayilari.get("İnceleniyor", 0),
            durum_sayilari.get("Çözüldü", 0),
        ],
    ]
    ozet_tablosu = Table(ozet_verileri, colWidths=[48 * mm] * 4)
    ozet_tablosu.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), ana_mavi),
                ("FONTNAME", (0, 1), (-1, 1), kalin_font),
                ("ALIGN", (0, 1), (-1, -1), "CENTER"),
                ("BOX", (0, 0), (-1, -1), 0.5, cizgi_rengi),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, cizgi_rengi),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    hikaye.extend(
        [
            Paragraph("Özet", bolum_stili),
            ozet_tablosu,
            Paragraph("İhlal Kayıtları", bolum_stili),
        ]
    )

    tablo_verileri = [
        [
            paragraf(baslik, tablo_baslik_stili)
            for baslik in (
                "ID",
                "Tarih",
                "İhlal Türü",
                "Güven Skoru",
                "Durum",
                "Açıklama",
            )
        ]
    ]
    for ihlal in ihlaller:
        guven_metni = (
            f"%{float(ihlal[3]):.2f}" if ihlal[3] is not None else "-"
        )
        tablo_verileri.append(
            [
                paragraf(ihlal[0]),
                paragraf(ihlal[1]),
                paragraf(ihlal[2]),
                paragraf(guven_metni),
                paragraf(ihlal[4]),
                paragraf(ihlal[5]),
            ]
        )

    ihlal_tablosu = Table(
        tablo_verileri,
        colWidths=[10 * mm, 32 * mm, 42 * mm, 22 * mm, 25 * mm, 126 * mm],
        repeatRows=1,
        splitByRow=1,
        splitInRow=1,
    )
    ihlal_tablosu.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), ana_mavi),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BOX", (0, 0), (-1, -1), 0.5, cizgi_rengi),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, cizgi_rengi),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F8FA")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    hikaye.extend([ihlal_tablosu, PageBreak()])

    dagilim_basligi = [
        paragraf("Tarih", tablo_baslik_stili),
        paragraf("İhlal Sayısı", tablo_baslik_stili),
    ]
    gunluk_veriler = [dagilim_basligi] + [
        [paragraf(tarih), adet]
        for tarih, adet in sorted(gunluk_sayilar.items())
    ]
    if len(gunluk_veriler) == 1:
        gunluk_veriler.append([paragraf("Kayıt yok"), 0])

    gunluk_tablosu = Table(
        gunluk_veriler,
        colWidths=[55 * mm, 35 * mm],
        repeatRows=1,
        hAlign="LEFT",
    )
    gunluk_tablosu.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), ana_mavi),
                ("FONTNAME", (1, 1), (1, -1), normal_font),
                ("ALIGN", (1, 1), (1, -1), "CENTER"),
                ("BOX", (0, 0), (-1, -1), 0.5, cizgi_rengi),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, cizgi_rengi),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )

    durum_verileri = [
        [
            paragraf("Durum", tablo_baslik_stili),
            paragraf("İhlal Sayısı", tablo_baslik_stili),
        ]
    ] + [
        [paragraf(durum), durum_sayilari.get(durum, 0)]
        for durum in IZINLI_IHLAL_DURUMLARI
    ]
    durum_tablosu = Table(
        durum_verileri,
        colWidths=[55 * mm, 35 * mm],
        repeatRows=1,
        hAlign="LEFT",
    )
    durum_tablosu.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), ana_mavi),
                ("FONTNAME", (1, 1), (1, -1), normal_font),
                ("ALIGN", (1, 1), (1, -1), "CENTER"),
                ("BOX", (0, 0), (-1, -1), 0.5, cizgi_rengi),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, cizgi_rengi),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    hikaye.extend(
        [
            Paragraph("Günlük Dağılım", bolum_stili),
            gunluk_tablosu,
            Paragraph("Durum Dağılımı", bolum_stili),
            durum_tablosu,
        ]
    )

    def sayfa_altligi(canvas, doc):
        canvas.saveState()
        canvas.setFont(normal_font, 7)
        canvas.setFillColor(colors.HexColor("#64748B"))
        canvas.drawString(10 * mm, 7 * mm, "İSG AI Sistemi")
        canvas.drawRightString(
            sayfa_boyutu[0] - 10 * mm,
            7 * mm,
            f"Sayfa {doc.page}",
        )
        canvas.restoreState()

    belge.build(
        hikaye,
        onFirstPage=sayfa_altligi,
        onLaterPages=sayfa_altligi,
    )
    output.seek(0)
    return output


@app.route("/raporlar")
@yetki_gerekli("raporlar")
def raporlar():
    baglanti = connect_database()
    ihlal_turleri, filtreler, ihlaller, hata = rapor_verilerini_getir(
        baglanti,
        request.args,
    )

    if hata:
        baglanti.close()
        return hata, 400

    kamera_listesi = baglanti.execute("SELECT id,kamera_adi FROM kameralar ORDER BY id").fetchall()
    baglanti.close()

    durum_sayilari = Counter(row[4] for row in ihlaller)
    gunluk_sayilar = Counter(row[1][:10] for row in ihlaller)
    gunluk_dagilim = sorted(gunluk_sayilar.items())
    durum_dagilim = [
        (durum, durum_sayilari.get(durum, 0))
        for durum in IZINLI_IHLAL_DURUMLARI
    ]

    return render_template(
        "raporlar.html",
        ihlaller=ihlaller,
        kameralar=[{"id": row[0], "kamera_adi": row[1]} for row in kamera_listesi],
        filtreler=filtreler,
        ihlal_turleri=ihlal_turleri,
        durumlar=IZINLI_IHLAL_DURUMLARI,
        toplam_ihlal=len(ihlaller),
        yeni_sayisi=durum_sayilari.get("Yeni", 0),
        inceleniyor_sayisi=durum_sayilari.get("İnceleniyor", 0),
        cozuldu_sayisi=durum_sayilari.get("Çözüldü", 0),
        gunluk_dagilim=gunluk_dagilim,
        durum_dagilim=durum_dagilim,
    )


@app.route("/raporlar/csv")
@yetki_gerekli("raporlar")
def raporlar_csv():
    baglanti = connect_database()

    try:
        _, filtreler, ihlaller, hata = rapor_verilerini_getir(
            baglanti,
            request.args,
            aciklama_dahil=True,
        )

        if hata:
            return hata, 400

        audit_kaydi_ekle(
            baglanti,
            "rapor_csv_disari_aktarildi",
            "rapor",
            None,
            yeni_deger={
                "filtreler": filtreler,
                "kayit_sayisi": len(ihlaller),
            },
        )
        baglanti.commit()
    except Exception:
        baglanti.rollback()
        raise
    finally:
        baglanti.close()

    output = io.StringIO(newline="")
    writer = csv.writer(output, delimiter=";", lineterminator="\r\n")
    writer.writerow(
        [
            "ID",
            "Tarih",
            "İhlal Türü",
            "Güven Skoru",
            "Durum",
            "Açıklama",
        ]
    )

    for ihlal in ihlaller:
        writer.writerow([csv_hucresini_guvenli_yap(value) for value in ihlal])

    dosya_adi = datetime.now().strftime("isg_raporu_%Y%m%d_%H%M%S.csv")
    csv_bytes = output.getvalue().encode("utf-8-sig")
    return Response(
        csv_bytes,
        content_type="text/csv; charset=utf-8",
        headers={
            "Content-Disposition": f'attachment; filename="{dosya_adi}"'
        },
    )


@app.route("/raporlar/xlsx")
@yetki_gerekli("raporlar")
def raporlar_xlsx():
    baglanti = connect_database()

    try:
        _, filtreler, ihlaller, hata = rapor_verilerini_getir(
            baglanti,
            request.args,
            aciklama_dahil=True,
        )

        if hata:
            return hata, 400

        xlsx_output = xlsx_workbook_olustur(ihlaller, filtreler)
        audit_kaydi_ekle(
            baglanti,
            "rapor_xlsx_disari_aktarildi",
            "rapor",
            None,
            yeni_deger={
                "filtreler": filtreler,
                "kayit_sayisi": len(ihlaller),
            },
        )
        baglanti.commit()
    except XlsxBagimliligiHatasi:
        baglanti.rollback()
        app.logger.exception("XLSX bağımlılığı yüklenemedi.")
        return "Excel dışa aktarma özelliği şu anda kullanılamıyor.", 503
    except Exception:
        baglanti.rollback()
        app.logger.exception("XLSX raporu oluşturulamadı.")
        return "Excel raporu oluşturulamadı. Lütfen tekrar deneyin.", 500
    finally:
        baglanti.close()

    dosya_adi = datetime.now().strftime(
        "isg_ihlal_raporu_%Y-%m-%d.xlsx"
    )
    return send_file(
        xlsx_output,
        as_attachment=True,
        download_name=dosya_adi,
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )


@app.route("/raporlar/pdf")
@yetki_gerekli("raporlar")
def raporlar_pdf():
    baglanti = connect_database()

    try:
        _, filtreler, ihlaller, hata = rapor_verilerini_getir(
            baglanti,
            request.args,
            aciklama_dahil=True,
        )

        if hata:
            return hata, 400

        pdf_output = pdf_raporu_olustur(ihlaller, filtreler)
        audit_kaydi_ekle(
            baglanti,
            "rapor_pdf_disari_aktarildi",
            "rapor",
            None,
            yeni_deger={
                "filtreler": filtreler,
                "kayit_sayisi": len(ihlaller),
            },
        )
        baglanti.commit()
    except (PdfBagimliligiHatasi, PdfFontHatasi):
        baglanti.rollback()
        app.logger.exception("PDF bağımlılığı veya Unicode font yüklenemedi.")
        return "PDF dışa aktarma özelliği şu anda kullanılamıyor.", 503
    except Exception:
        baglanti.rollback()
        app.logger.exception("PDF raporu oluşturulamadı.")
        return "PDF raporu oluşturulamadı. Lütfen tekrar deneyin.", 500
    finally:
        baglanti.close()

    dosya_adi = datetime.now().strftime(
        "isg_ihlal_raporu_%Y-%m-%d.pdf"
    )
    return send_file(
        pdf_output,
        as_attachment=True,
        download_name=dosya_adi,
        mimetype="application/pdf",
    )


# --------------------------------------------------
# KAMERA YÖNETİMİ
# --------------------------------------------------

@app.route("/kameralar")
@yetki_gerekli("kameralar")
def kameralar():
    baglanti = connect_database()
    kamera_listesi = baglanti.execute(
        """
        SELECT
            id,
            kamera_adi,
            kamera_index,
            aktif,
            yasak_bolge_orani,
            olusturma_tarihi,
            guncelleme_tarihi
        FROM kameralar
        ORDER BY id
        """
    ).fetchall()
    baglanti.close()

    return render_template("kameralar.html", kameralar=kamera_listesi)


@app.route("/kameralar/ekle", methods=["POST"])
@yetki_gerekli("kameralar")
@admin_gerekli
def kamera_ekle():
    kamera_degerleri, hata = kamera_formunu_dogrula(request.form)

    if hata:
        return hata, 400

    simdi = datetime.now().isoformat(timespec="seconds")
    baglanti = connect_database()

    try:
        baglanti.execute("BEGIN IMMEDIATE")
        duplicate = baglanti.execute(
            "SELECT id FROM kameralar WHERE kamera_index = ?",
            (kamera_degerleri["kamera_index"],),
        ).fetchone()

        if duplicate is not None:
            baglanti.rollback()
            return "Bu kamera indexi başka bir kamera tarafından kullanılıyor.", 400

        onceki_aktif_kameralar = baglanti.execute(
            "SELECT id FROM kameralar WHERE aktif = 1 ORDER BY id"
        ).fetchall()
        baglanti.execute(
            """
            UPDATE kameralar
            SET aktif = 0, guncelleme_tarihi = ?
            WHERE aktif = 1
            """,
            (simdi,),
        )
        cursor = baglanti.execute(
            """
            INSERT INTO kameralar (
                kamera_adi,
                kamera_index,
                aktif,
                yasak_bolge_orani,
                olusturma_tarihi,
                guncelleme_tarihi
            )
            VALUES (?, ?, 1, ?, ?, ?)
            """,
            (
                kamera_degerleri["kamera_adi"],
                kamera_degerleri["kamera_index"],
                kamera_degerleri["yasak_bolge_orani"],
                simdi,
                simdi,
            ),
        )
        kamera_id = cursor.lastrowid

        for onceki_kamera in onceki_aktif_kameralar:
            audit_kaydi_ekle(
                baglanti,
                "kamera_aktifligi_degistirildi",
                "kamera",
                onceki_kamera[0],
                eski_deger={"aktif": True},
                yeni_deger={
                    "aktif": False,
                    "otomatik": True,
                    "yeni_aktif_kamera_id": kamera_id,
                },
            )

        audit_kaydi_ekle(
            baglanti,
            "kamera_olusturuldu",
            "kamera",
            kamera_id,
            yeni_deger={
                **kamera_degerleri,
                "aktif": True,
                "onceki_aktif_kamera_idleri": [
                    row[0] for row in onceki_aktif_kameralar
                ],
            },
        )
        baglanti.commit()
    except sqlite3.IntegrityError:
        baglanti.rollback()
        return "Kamera kaydı mevcut kısıtlarla çakışıyor.", 400
    except Exception:
        baglanti.rollback()
        raise
    finally:
        baglanti.close()

    return redirect(url_for("kameralar"))


@app.route("/kameralar/<int:kamera_id>/duzenle", methods=["POST"])
@yetki_gerekli("kameralar")
@admin_gerekli
def kamera_duzenle(kamera_id):
    kamera_degerleri, hata = kamera_formunu_dogrula(request.form)

    if hata:
        return hata, 400

    baglanti = connect_database()

    try:
        baglanti.execute("BEGIN IMMEDIATE")
        kamera = baglanti.execute(
            """
            SELECT kamera_adi, kamera_index, aktif, yasak_bolge_orani
            FROM kameralar
            WHERE id = ?
            """,
            (kamera_id,),
        ).fetchone()

        if kamera is None:
            baglanti.rollback()
            return "Kamera bulunamadı.", 404

        duplicate = baglanti.execute(
            """
            SELECT id
            FROM kameralar
            WHERE kamera_index = ? AND id != ?
            """,
            (kamera_degerleri["kamera_index"], kamera_id),
        ).fetchone()

        if duplicate is not None:
            baglanti.rollback()
            return "Bu kamera indexi başka bir kamera tarafından kullanılıyor.", 400

        simdi = datetime.now().isoformat(timespec="seconds")
        baglanti.execute(
            """
            UPDATE kameralar
            SET kamera_adi = ?,
                kamera_index = ?,
                yasak_bolge_orani = ?,
                guncelleme_tarihi = ?
            WHERE id = ?
            """,
            (
                kamera_degerleri["kamera_adi"],
                kamera_degerleri["kamera_index"],
                kamera_degerleri["yasak_bolge_orani"],
                simdi,
                kamera_id,
            ),
        )
        audit_kaydi_ekle(
            baglanti,
            "kamera_ayarlari_degistirildi",
            "kamera",
            kamera_id,
            eski_deger={
                "kamera_adi": kamera[0],
                "kamera_index": kamera[1],
                "yasak_bolge_orani": kamera[3],
            },
            yeni_deger=kamera_degerleri,
        )
        baglanti.commit()
    except sqlite3.IntegrityError:
        baglanti.rollback()
        return "Kamera ayarları mevcut kısıtlarla çakışıyor.", 400
    except Exception:
        baglanti.rollback()
        raise
    finally:
        baglanti.close()

    return redirect(url_for("kameralar"))


@app.route("/kameralar/<int:kamera_id>/durum", methods=["POST"])
@yetki_gerekli("kameralar")
@admin_gerekli
def kamera_durum(kamera_id):
    baglanti = connect_database()

    try:
        baglanti.execute("BEGIN IMMEDIATE")
        kamera = baglanti.execute(
            "SELECT aktif FROM kameralar WHERE id = ?",
            (kamera_id,),
        ).fetchone()

        if kamera is None:
            baglanti.rollback()
            return "Kamera bulunamadı.", 404

        if kamera[0] == 1:
            baglanti.rollback()
            return (
                "Tek aktif kamera doğrudan pasifleştirilemez. "
                "Önce başka bir kamerayı aktif hale getirin.",
                400,
            )

        simdi = datetime.now().isoformat(timespec="seconds")
        onceki_aktif_kameralar = baglanti.execute(
            "SELECT id FROM kameralar WHERE aktif = 1 ORDER BY id"
        ).fetchall()
        baglanti.execute(
            """
            UPDATE kameralar
            SET aktif = 0, guncelleme_tarihi = ?
            WHERE aktif = 1
            """,
            (simdi,),
        )
        baglanti.execute(
            """
            UPDATE kameralar
            SET aktif = 1, guncelleme_tarihi = ?
            WHERE id = ?
            """,
            (
                simdi,
                kamera_id,
            ),
        )

        for onceki_kamera in onceki_aktif_kameralar:
            audit_kaydi_ekle(
                baglanti,
                "kamera_aktifligi_degistirildi",
                "kamera",
                onceki_kamera[0],
                eski_deger={"aktif": True},
                yeni_deger={
                    "aktif": False,
                    "otomatik": True,
                    "yeni_aktif_kamera_id": kamera_id,
                },
            )

        audit_kaydi_ekle(
            baglanti,
            "kamera_aktifligi_degistirildi",
            "kamera",
            kamera_id,
            eski_deger={"aktif": False},
            yeni_deger={
                "aktif": True,
                "onceki_aktif_kamera_idleri": [
                    row[0] for row in onceki_aktif_kameralar
                ],
            },
        )
        baglanti.commit()
    except sqlite3.IntegrityError:
        baglanti.rollback()
        return "Kamera aktivasyonu mevcut kısıtlarla çakışıyor.", 400
    except Exception:
        baglanti.rollback()
        raise
    finally:
        baglanti.close()

    return redirect(url_for("kameralar"))


# --------------------------------------------------
# KULLANICI YÖNETİMİ
# --------------------------------------------------

@app.route("/kullanicilar")
@yetki_gerekli("kullanicilar")
def kullanicilar():
    baglanti = connect_database()
    cursor = baglanti.cursor()

    cursor.execute(
        """
        SELECT
            id,
            ad_soyad,
            kullanici_adi,
            rol,
            aktif, eposta
        FROM kullanicilar
        ORDER BY id DESC
        """
    )

    kullanici_listesi = cursor.fetchall()
    baglanti.close()

    return render_template(
        "kullanicilar.html",
        kullanicilar=kullanici_listesi,
        izinli_roller=IZINLI_ROLLER,
    )


@app.route("/kullanicilar/ekle", methods=["POST"])
@yetki_gerekli("kullanicilar")
@admin_gerekli
def kullanici_ekle():
    ad_soyad = request.form.get("ad_soyad", "").strip()
    kullanici_adi = request.form.get("kullanici_adi", "").strip()
    sifre = request.form.get("sifre", "")
    rol = request.form.get("rol", "Görüntüleyici")
    eposta = request.form.get("eposta", "").strip()
    if eposta and (len(eposta) > 254 or "@" not in eposta or any(c.isspace() for c in eposta) or kontrol_karakteri_iceriyor(eposta)):
        return "Geçersiz e-posta.", 400

    if not metin_gecerli(
        ad_soyad,
        AD_SOYAD_MIN_UZUNLUK,
        AD_SOYAD_MAX_UZUNLUK,
    ):
        return "Ad soyad alanı geçersiz.", 400

    if not metin_gecerli(
        kullanici_adi,
        KULLANICI_ADI_MIN_UZUNLUK,
        KULLANICI_ADI_MAX_UZUNLUK,
    ):
        return "Kullanıcı adı geçersiz.", 400

    if not sifre or len(sifre) > PAROLA_MAX_UZUNLUK:
        return "Parola alanı geçersiz.", 400

    if rol not in IZINLI_ROLLER:
        return "Geçersiz rol.", 400

    is_valid, error_message = validate_password(sifre)

    if not is_valid:
        flash(error_message, "error")
        return redirect(url_for("kullanicilar"))

    sifre_hash = generate_password_hash(sifre)

    baglanti = connect_database()
    cursor = baglanti.cursor()

    try:
        cursor.execute(
            """
            INSERT INTO kullanicilar
                (ad_soyad, kullanici_adi, sifre, rol, aktif, eposta)
            VALUES (?, ?, ?, ?, 1, ?)
            """,
            (ad_soyad, kullanici_adi, sifre_hash, rol, eposta or None),
        )
        yeni_kullanici_id = cursor.lastrowid
        audit_kaydi_ekle(
            baglanti,
            "kullanici_olusturuldu",
            "kullanici",
            yeni_kullanici_id,
            yeni_deger={
                "ad_soyad": ad_soyad,
                "kullanici_adi": kullanici_adi,
                "rol": rol,
                "aktif": True,
            },
        )
        baglanti.commit()

    except sqlite3.IntegrityError:
        baglanti.rollback()
        flash("Bu kullanıcı adı zaten kullanılıyor.", "error")
        return redirect(url_for("kullanicilar"))
    except Exception:
        baglanti.rollback()
        raise
    finally:
        baglanti.close()

    return redirect(url_for("kullanicilar"))


@app.route("/kullanicilar/<int:kullanici_id>/durum", methods=["POST"])
@yetki_gerekli("kullanicilar")
@admin_gerekli
def kullanici_durum(kullanici_id):
    if kullanici_id == session.get("kullanici_id"):
        return "Kendi hesabınızı pasif yapamazsınız.", 400

    baglanti = connect_database()

    try:
        baglanti.execute("BEGIN IMMEDIATE")
        kullanici = baglanti.execute(
            """
            SELECT rol, aktif
            FROM kullanicilar
            WHERE id = ?
            """,
            (kullanici_id,),
        ).fetchone()

        if kullanici is None:
            baglanti.rollback()
            return "Kullanıcı bulunamadı.", 404

        yeni_durum = 0 if kullanici[1] == 1 else 1

        if (
            kullanici[0] == ADMIN_ROLU
            and kullanici[1] == 1
            and yeni_durum == 0
            and aktif_admin_sayisi(baglanti) <= 1
        ):
            baglanti.rollback()
            return "Son aktif yönetici pasif yapılamaz.", 400

        baglanti.execute(
            """
            UPDATE kullanicilar
            SET aktif = ?
            WHERE id = ?
            """,
            (yeni_durum, kullanici_id),
        )
        audit_kaydi_ekle(
            baglanti,
            "kullanici_aktifligi_degistirildi",
            "kullanici",
            kullanici_id,
            eski_deger={"aktif": bool(kullanici[1])},
            yeni_deger={"aktif": bool(yeni_durum)},
        )
        baglanti.commit()
    except Exception:
        baglanti.rollback()
        raise
    finally:
        baglanti.close()

    return redirect(url_for("kullanicilar"))


@app.route("/kullanicilar/<int:kullanici_id>/rol", methods=["POST"])
@yetki_gerekli("kullanicilar")
@admin_gerekli
def kullanici_rol(kullanici_id):
    yeni_rol = request.form.get("rol")

    if yeni_rol not in IZINLI_ROLLER:
        return "Geçersiz rol.", 400

    if kullanici_id == session.get("kullanici_id"):
        return "Kendi rolünüzü buradan değiştiremezsiniz.", 400

    baglanti = connect_database()

    try:
        baglanti.execute("BEGIN IMMEDIATE")
        kullanici = baglanti.execute(
            """
            SELECT rol, aktif
            FROM kullanicilar
            WHERE id = ?
            """,
            (kullanici_id,),
        ).fetchone()

        if kullanici is None:
            baglanti.rollback()
            return "Kullanıcı bulunamadı.", 404

        if (
            kullanici[0] == ADMIN_ROLU
            and kullanici[1] == 1
            and yeni_rol != ADMIN_ROLU
            and aktif_admin_sayisi(baglanti) <= 1
        ):
            baglanti.rollback()
            return "Son aktif yöneticinin rolü değiştirilemez.", 400

        baglanti.execute(
            """
            UPDATE kullanicilar
            SET rol = ?
            WHERE id = ?
            """,
            (yeni_rol, kullanici_id),
        )
        audit_kaydi_ekle(
            baglanti,
            "kullanici_rolu_degistirildi",
            "kullanici",
            kullanici_id,
            eski_deger={"rol": kullanici[0]},
            yeni_deger={"rol": yeni_rol},
        )
        baglanti.commit()
    except Exception:
        baglanti.rollback()
        raise
    finally:
        baglanti.close()

    return redirect(url_for("kullanicilar"))


@app.route(
    "/kullanicilar/<int:kullanici_id>/yetkiler",
    methods=["GET", "POST"],
)
@yetki_gerekli("kullanicilar")
@admin_gerekli
def kullanici_yetkileri(kullanici_id):
    baglanti = connect_database()
    cursor = baglanti.cursor()

    cursor.execute(
        """
        SELECT
            id,
            ad_soyad,
            kullanici_adi
        FROM kullanicilar
        WHERE id = ?
        """,
        (kullanici_id,),
    )

    kullanici = cursor.fetchone()

    if kullanici is None:
        baglanti.close()
        return "Kullanıcı bulunamadı.", 404

    izinli_moduller = [modul_kodu for modul_kodu, _ in MODULLER]

    if request.method == "POST":
        secilen_yetkiler = request.form.getlist("yetkiler")

        if any(modul not in izinli_moduller for modul in secilen_yetkiler):
            baglanti.close()
            return "Geçersiz modül yetkisi.", 400

        secilen_yetkiler = [
            modul for modul in secilen_yetkiler if modul in izinli_moduller
        ]

        if (
            kullanici_id == session.get("kullanici_id")
            and session.get("rol") == ADMIN_ROLU
            and "kullanicilar" not in secilen_yetkiler
        ):
            baglanti.close()
            return "Kendi kullanıcı yönetimi yetkinizi kaldıramazsınız.", 400

        try:
            baglanti.execute("BEGIN IMMEDIATE")
            onceki_yetkiler = [
                kayit[0]
                for kayit in baglanti.execute(
                    """
                    SELECT modul
                    FROM kullanici_yetkileri
                    WHERE kullanici_id = ?
                    ORDER BY modul
                    """,
                    (kullanici_id,),
                ).fetchall()
            ]
            yeni_yetkiler = sorted(set(secilen_yetkiler))

            baglanti.execute(
                """
                DELETE FROM kullanici_yetkileri
                WHERE kullanici_id = ?
                """,
                (kullanici_id,),
            )

            for modul in yeni_yetkiler:
                baglanti.execute(
                    """
                    INSERT INTO kullanici_yetkileri
                        (kullanici_id, modul)
                    VALUES (?, ?)
                    """,
                    (kullanici_id, modul),
                )

            audit_kaydi_ekle(
                baglanti,
                "kullanici_yetkileri_degistirildi",
                "kullanici",
                kullanici_id,
                eski_deger={"yetkiler": onceki_yetkiler},
                yeni_deger={
                    "yetkiler": yeni_yetkiler,
                    "eklenen": sorted(set(yeni_yetkiler) - set(onceki_yetkiler)),
                    "kaldirilan": sorted(
                        set(onceki_yetkiler) - set(yeni_yetkiler)
                    ),
                },
            )
            baglanti.commit()
        except Exception:
            baglanti.rollback()
            baglanti.close()
            raise

        if kullanici_id == session.get("kullanici_id"):
            session["yetkiler"] = yeni_yetkiler

        baglanti.close()

        # Kendi Kullanıcılar yetkisini kaldırdıysa kullanıcılar sayfasına
        # yönlendirmeye çalışmayalım.
        if kullanici_id == session.get("kullanici_id"):
            return redirect(ilk_erisilebilir_sayfa())

        return redirect(url_for("kullanicilar"))

    cursor.execute(
        """
        SELECT modul
        FROM kullanici_yetkileri
        WHERE kullanici_id = ?
        """,
        (kullanici_id,),
    )

    mevcut_yetkiler = [kayit[0] for kayit in cursor.fetchall()]
    baglanti.close()

    return render_template(
        "yetkiler.html",
        kullanici=kullanici,
        moduller=MODULLER,
        mevcut_yetkiler=mevcut_yetkiler,
    )


# --------------------------------------------------
# UYGULAMA
# --------------------------------------------------

from web_app.panel import register_panel
register_panel(app, sys.modules[__name__])


if __name__ == "__main__":
    debug_mode = os.environ.get("ISG_FLASK_DEBUG", "0") == "1"
    app.run(debug=debug_mode)
